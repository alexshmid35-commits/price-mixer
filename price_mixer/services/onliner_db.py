"""Local Onliner catalog SQLite store and import helpers."""

from __future__ import annotations

from collections.abc import Callable
from contextlib import contextmanager
import csv
import os
from pathlib import Path
import re
import shutil
import sqlite3
import tempfile
import threading
import time
import urllib.error
import urllib.parse
import urllib.request

from price_mixer.logging_config import get_logger, log_context, new_job_id
from price_mixer.runtime_paths import get_runtime_paths
from price_mixer.services.product_normalization import (
    infer_category,
    normalize_catalog_category_name,
    normalize_onliner_id,
)

LOGGER = get_logger("price_mixer.database.catalog")

RUNTIME_PATHS = get_runtime_paths()
ONLINER_DB_FILE = RUNTIME_PATHS.data_file("onliner_products.db")
ONLINER_DB_GSHEET_CACHE_DIR = RUNTIME_PATHS.cache_dir / "onliner_db_gsheet"
ONLINER_DB_GSHEET_CACHE_TTL_SEC = 12 * 60 * 60

DB_WRITE_LOCK = threading.RLock()
CATALOG_IMPORT_LOCK = threading.RLock()
catalog_import_status = {
    "running": False,
    "total": 0,
    "done": 0,
    "inserted": 0,
    "skipped": 0,
    "message": "",
    "percent": 0,
    "finished_at": None,
}


def _is_corrupt_db_error(exc):
    text = str(exc or "").lower()
    return "database disk image is malformed" in text or "file is not a database" in text


def quarantine_onliner_db(db_file=None):
    db_file = Path(db_file or ONLINER_DB_FILE)
    stamp = time.strftime("%Y%m%d-%H%M%S")
    moved = []
    for suffix in ("", "-wal", "-shm"):
        path = Path(str(db_file) + suffix)
        if not path.exists():
            continue
        backup = Path(str(path) + f".corrupt-{stamp}")
        try:
            os.replace(path, backup)
            moved.append(str(backup))
        except OSError:
            pass
    if moved:
        LOGGER.error("corrupt catalog database quarantined files=%s", len(moved))
    return moved


def db_connect(db_file=None):
    db_file = db_file or ONLINER_DB_FILE
    conn = sqlite3.connect(str(db_file), check_same_thread=False, timeout=15)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    conn.row_factory = sqlite3.Row
    return conn


@contextmanager
def db_connection(db_file=None):
    conn = db_connect(db_file)
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def catalog_revision(db_file=None):
    """Return a stable revision token for candidate-cache invalidation."""
    try:
        with db_connection(db_file) as conn:
            catalog = conn.execute(
                "SELECT COUNT(*),COALESCE(MAX(updated_at),0) FROM onliner_catalog"
            ).fetchone()
            names = conn.execute(
                "SELECT COUNT(*),COALESCE(MAX(updated_at),0) FROM name_index"
            ).fetchone()
        return ":".join(str(int(value or 0)) for value in (*catalog, *names))
    except Exception:
        return "unavailable"


def _ensure_catalog_category_column(conn):
    columns = {str(row[1]) for row in conn.execute("PRAGMA table_info(onliner_catalog)").fetchall()}
    if "category" not in columns:
        conn.execute("ALTER TABLE onliner_catalog ADD COLUMN category TEXT DEFAULT ''")


def init_onliner_db(db_file=None):
    db_file = db_file or ONLINER_DB_FILE
    try:
        with db_connection(db_file) as conn:
            conn.executescript("""
                CREATE TABLE IF NOT EXISTS onliner_catalog (
                    onliner_id  TEXT PRIMARY KEY,
                    name        TEXT NOT NULL,
                    url         TEXT DEFAULT '',
                    category    TEXT DEFAULT '',
                    source      TEXT DEFAULT '',
                    updated_at  INTEGER DEFAULT 0
                );
                CREATE TABLE IF NOT EXISTS name_index (
                    name_key    TEXT PRIMARY KEY,
                    onliner_id  TEXT NOT NULL,
                    raw_name    TEXT NOT NULL,
                    source      TEXT DEFAULT '',
                    updated_at  INTEGER DEFAULT 0
                );
                CREATE INDEX IF NOT EXISTS idx_ni_oid
                    ON name_index(onliner_id);
                CREATE INDEX IF NOT EXISTS idx_ni_rawname
                    ON name_index(raw_name);
            """)
            _ensure_catalog_category_column(conn)
            check = str(conn.execute("PRAGMA quick_check").fetchone()[0] or "")
            if check.lower() != "ok":
                raise sqlite3.DatabaseError(check)
        LOGGER.info("catalog database initialized")
    except Exception as exc:
        if _is_corrupt_db_error(exc) or "invalid page" in str(exc).lower():
            quarantine_onliner_db(db_file)
            try:
                with db_connection(db_file) as conn:
                    conn.executescript("""
                        CREATE TABLE IF NOT EXISTS onliner_catalog (
                            onliner_id  TEXT PRIMARY KEY,
                            name        TEXT NOT NULL,
                            url         TEXT DEFAULT '',
                            category    TEXT DEFAULT '',
                            source      TEXT DEFAULT '',
                            updated_at  INTEGER DEFAULT 0
                        );
                        CREATE TABLE IF NOT EXISTS name_index (
                            name_key    TEXT PRIMARY KEY,
                            onliner_id  TEXT NOT NULL,
                            raw_name    TEXT NOT NULL,
                            source      TEXT DEFAULT '',
                            updated_at  INTEGER DEFAULT 0
                        );
                        CREATE INDEX IF NOT EXISTS idx_ni_oid
                            ON name_index(onliner_id);
                        CREATE INDEX IF NOT EXISTS idx_ni_rawname
                            ON name_index(raw_name);
                    """)
                    _ensure_catalog_category_column(conn)
                LOGGER.warning("catalog database recreated after corruption")
                return
            except Exception as recreate_exc:
                LOGGER.exception("catalog database recreation failed")
        LOGGER.error("catalog database initialization failed: %s", exc)


def populate_from_df(df, source_label, *, normalize_name_key, skip_suppliers=None):
    if df is None or df.empty:
        return 0, 0
    skip_upper = {str(s).upper() for s in (skip_suppliers or [])}
    supplier_col = "Поставщик" if "Поставщик" in df.columns else None
    now = int(time.time())

    products = {}
    names = {}
    for _, row in df.iterrows():
        oid = normalize_onliner_id(row.get("OnlinerID", ""))
        if not oid:
            continue
        name = str(row.get("Название", "")).strip()
        if not name:
            continue
        if supplier_col:
            supplier = str(row.get(supplier_col, "")).strip().upper()
            if supplier in skip_upper:
                continue
        url = str(row.get("Ссылка", "")).strip()
        name_key = normalize_name_key(name)
        if oid not in products:
            products[oid] = (name, url)
        if name_key and name_key not in names:
            names[name_key] = (oid, name)

    if not products:
        return 0, 0

    try:
        with DB_WRITE_LOCK:
            with db_connection() as conn:
                conn.executemany(
                    "INSERT INTO onliner_catalog"
                    "(onliner_id, name, url, source, updated_at) VALUES(?,?,?,?,?) "
                    "ON CONFLICT(onliner_id) DO UPDATE SET "
                    "name=excluded.name, url=excluded.url, source=excluded.source, updated_at=excluded.updated_at",
                    [(oid, name, url, source_label, now) for oid, (name, url) in products.items()],
                )
                conn.executemany(
                    "INSERT OR REPLACE INTO name_index"
                    "(name_key, onliner_id, raw_name, source, updated_at) VALUES(?,?,?,?,?)",
                    [(key, oid, raw, source_label, now) for key, (oid, raw) in names.items()],
                )
        LOGGER.info(
            "catalog populated source=%s products=%s names=%s",
            source_label,
            len(products),
            len(names),
        )
        return len(products), len(names)
    except Exception as exc:
        LOGGER.warning("catalog population write failed: %s", exc)
        return 0, 0


def upsert_product(onliner_id, name, url, *, normalize_name_key, source="manual"):
    oid = normalize_onliner_id(onliner_id)
    name = str(name or "").strip()
    if not oid or not name:
        return
    name_key = normalize_name_key(name)
    now = int(time.time())
    try:
        with DB_WRITE_LOCK:
            with db_connection() as conn:
                conn.execute(
                    "INSERT INTO onliner_catalog"
                    "(onliner_id, name, url, source, updated_at) VALUES(?,?,?,?,?) "
                    "ON CONFLICT(onliner_id) DO UPDATE SET "
                    "name=excluded.name, url=excluded.url, source=excluded.source, updated_at=excluded.updated_at",
                    (oid, name, url or "", source, now),
                )
                if name_key:
                    conn.execute(
                        "INSERT OR REPLACE INTO name_index"
                        "(name_key, onliner_id, raw_name, source, updated_at) VALUES(?,?,?,?,?)",
                        (name_key, oid, name, source, now),
                    )
    except Exception as exc:
        LOGGER.warning("catalog product upsert failed: %s", exc)


def get_product_by_id(onliner_id):
    oid = normalize_onliner_id(onliner_id)
    if not oid:
        return None
    try:
        with db_connection() as conn:
            row = conn.execute(
                "SELECT onliner_id, name, url, category, source, updated_at "
                "FROM onliner_catalog WHERE onliner_id = ? LIMIT 1",
                (oid,),
            ).fetchone()
        if not row:
            return None
        return {
            "id": normalize_onliner_id(row["onliner_id"]),
            "name": str(row["name"] or "").strip(),
            "url": str(row["url"] or "").strip(),
            "category": normalize_catalog_category_name(str(row["category"] or "").strip()),
            "source": str(row["source"] or "").strip(),
            "updated_at": int(row["updated_at"] or 0),
        }
    except Exception as exc:
        LOGGER.warning("catalog product lookup failed: %s", exc)
        return None


def get_categories_by_ids(onliner_ids):
    ids = []
    seen = set()
    for value in onliner_ids or []:
        oid = normalize_onliner_id(value)
        if oid and oid not in seen:
            seen.add(oid)
            ids.append(oid)
    if not ids:
        return {}
    result = {}
    try:
        with db_connection() as conn:
            for start in range(0, len(ids), 800):
                batch = ids[start:start + 800]
                placeholders = ",".join("?" for _ in batch)
                rows = conn.execute(
                    "SELECT onliner_id, category FROM onliner_catalog "
                    f"WHERE onliner_id IN ({placeholders})",
                    batch,
                ).fetchall()
                for row in rows:
                    oid = normalize_onliner_id(row["onliner_id"])
                    category = normalize_catalog_category_name(str(row["category"] or "").strip())
                    if oid and category:
                        result[oid] = category
    except Exception as exc:
        LOGGER.warning("catalog category lookup by IDs failed: %s", exc)
    return result


def get_distinct_categories():
    categories = set()
    try:
        with db_connection() as conn:
            rows = conn.execute(
                "SELECT DISTINCT category FROM onliner_catalog "
                "WHERE TRIM(COALESCE(category, '')) <> ''"
            ).fetchall()
            for row in rows:
                category = normalize_catalog_category_name(str(row["category"] or "").strip())
                if category:
                    categories.add(category)
    except Exception as exc:
        LOGGER.warning("catalog distinct categories lookup failed: %s", exc)
    return sorted(categories)


def update_categories(items, source="onliner_parser"):
    rows = []
    now = int(time.time())
    for item in items or []:
        oid = normalize_onliner_id((item or {}).get("onliner_id"))
        item_source = str((item or {}).get("source") or "").strip()
        name = str((item or {}).get("name") or "").strip() if item_source == "catalog_api_id" else ""
        category = normalize_catalog_category_name(str((item or {}).get("category") or "").strip())
        url = str((item or {}).get("url") or "").strip()
        if oid and category:
            rows.append((oid, name or oid, url, category, source, now))
    if not rows:
        return 0
    try:
        with DB_WRITE_LOCK:
            with db_connection() as conn:
                conn.executemany(
                    "INSERT INTO onliner_catalog "
                    "(onliner_id, name, url, category, source, updated_at) VALUES (?,?,?,?,?,?) "
                    "ON CONFLICT(onliner_id) DO UPDATE SET "
                    "name=CASE WHEN excluded.name <> excluded.onliner_id THEN excluded.name ELSE onliner_catalog.name END, "
                    "url=CASE WHEN excluded.url <> '' THEN excluded.url ELSE onliner_catalog.url END, "
                    "category=excluded.category, source=excluded.source, updated_at=excluded.updated_at",
                    rows,
                )
        return len(rows)
    except Exception as exc:
        LOGGER.warning("catalog category update failed: %s", exc)
        return 0


def find_exact_id_for_name(product_name, *, normalize_name_key):
    name = str(product_name or "").strip()
    if not name:
        return None
    name_key = normalize_name_key(name)
    if not name_key:
        return None
    try:
        with db_connection() as conn:
            row = conn.execute(
                "SELECT ni.onliner_id, ni.raw_name, oc.url "
                "FROM name_index ni "
                "LEFT JOIN onliner_catalog oc ON oc.onliner_id = ni.onliner_id "
                "WHERE ni.name_key = ? LIMIT 1",
                (name_key,),
            ).fetchone()
        if not row:
            return None
        return {
            "id": normalize_onliner_id(row["onliner_id"]),
            "name": str(row["raw_name"] or "").strip(),
            "url": str(row["url"] or "").strip(),
            "score": 1.0,
            "source": "db_exact",
        }
    except Exception as exc:
        LOGGER.warning("catalog exact ID lookup failed: %s", exc)
        return None


def _add_candidate_rows(target, rows):
    for row in rows:
        if row[0] not in target:
            target[row[0]] = (row[1], row[2])


def _best_candidate_match(calc_name_match, local_name, candidate_name, candidate_url=""):
    comparisons = [calc_name_match(local_name, candidate_name)]
    if str(candidate_url or "").strip():
        comparisons.append(calc_name_match(
            local_name,
            f"{candidate_name} {candidate_url}".strip(),
        ))
    return max(
        comparisons,
        key=lambda item: (bool((item or {}).get("match", False)), float((item or {}).get("score", 0.0) or 0.0)),
    )


def find_id_for_name(
    product_name,
    *,
    normalize_name_key,
    raw_search_tokens,
    model_hint_tokens,
    preferred_brand_token,
    strict_candidate_allowed,
    calc_name_match,
    b2b_search_candidates=None,
    threshold=0.75,
    allow_b2b=True,
):
    name = str(product_name or "").strip()
    if not name:
        return None
    name_key = normalize_name_key(name)

    try:
        with db_connection() as conn:
            if name_key:
                row = conn.execute(
                    "SELECT ni.onliner_id, ni.raw_name, oc.url "
                    "FROM name_index ni "
                    "JOIN onliner_catalog oc ON oc.onliner_id = ni.onliner_id "
                    "WHERE ni.name_key = ? LIMIT 1",
                    (name_key,),
                ).fetchone()
                if row:
                    return {"id": row[0], "name": row[1], "url": row[2], "score": 1.0, "source": "db_exact"}

            candidates = {}
            for token in raw_search_tokens(name):
                if len(token) < 4:
                    continue
                _add_candidate_rows(candidates, conn.execute(
                    "SELECT ni.onliner_id, ni.raw_name, oc.url "
                    "FROM name_index ni "
                    "JOIN onliner_catalog oc ON oc.onliner_id = ni.onliner_id "
                    "WHERE ni.raw_name LIKE ? LIMIT 80",
                    (f"%{token}%",),
                ).fetchall())
                number_match = re.match(r"^([A-Za-z]{1,4}[-]?\d{3,5})", token)
                if number_match and number_match.group(1) != token:
                    _add_candidate_rows(candidates, conn.execute(
                        "SELECT ni.onliner_id, ni.raw_name, oc.url "
                        "FROM name_index ni "
                        "JOIN onliner_catalog oc ON oc.onliner_id = ni.onliner_id "
                        "WHERE ni.raw_name LIKE ? LIMIT 40",
                        (f"%{number_match.group(1)}%",),
                    ).fetchall())

            if len(candidates) < 3:
                for token in model_hint_tokens(name)[:3]:
                    if len(token) < 4:
                        continue
                    _add_candidate_rows(candidates, conn.execute(
                        "SELECT ni.onliner_id, ni.raw_name, oc.url "
                        "FROM name_index ni "
                        "JOIN onliner_catalog oc ON oc.onliner_id = ni.onliner_id "
                        "WHERE ni.raw_name LIKE ? LIMIT 60",
                        (f"%{token}%",),
                    ).fetchall())

            if len(candidates) < 5:
                brand = preferred_brand_token(name)
                if brand and len(brand) >= 3:
                    _add_candidate_rows(candidates, conn.execute(
                        "SELECT ni.onliner_id, ni.raw_name, oc.url "
                        "FROM name_index ni "
                        "JOIN onliner_catalog oc ON oc.onliner_id = ni.onliner_id "
                        "WHERE ni.raw_name LIKE ? LIMIT 150",
                        (f"%{brand}%",),
                    ).fetchall())

            if len(candidates) < 5:
                numbers = re.findall(r"\b(\d{3,5})\b", name)
                brand = preferred_brand_token(name)
                for number in numbers[:2]:
                    if brand:
                        _add_candidate_rows(candidates, conn.execute(
                            "SELECT ni.onliner_id, ni.raw_name, oc.url "
                            "FROM name_index ni "
                            "JOIN onliner_catalog oc ON oc.onliner_id = ni.onliner_id "
                            "WHERE ni.raw_name LIKE ? AND ni.raw_name LIKE ? LIMIT 80",
                            (f"%{brand}%", f"%{number}%"),
                        ).fetchall())

        if not candidates:
            return None

        best_score = 0.0
        best = None
        for oid, (candidate_name, candidate_url) in candidates.items():
            allowed, _reason = strict_candidate_allowed(name, candidate_name)
            if not allowed:
                continue
            match = _best_candidate_match(calc_name_match, name, candidate_name, candidate_url)
            score = float(match.get("score", 0.0) or 0.0)
            if score > best_score:
                best_score = score
                best = {"id": oid, "name": candidate_name, "url": candidate_url}

        if best_score >= threshold and best:
            return {**best, "score": round(best_score, 3), "source": "db_fuzzy"}

        if allow_b2b and b2b_search_candidates:
            for candidate in b2b_search_candidates(name, category_name=infer_category(name), limit=8):
                score = float(candidate.get("score", 0.0) or 0.0)
                if score >= float(threshold):
                    return {
                        "id": normalize_onliner_id(candidate.get("id", "")),
                        "name": str(candidate.get("name", "") or "").strip(),
                        "url": str(candidate.get("url", "") or "").strip(),
                        "score": round(score, 3),
                        "source": "b2b_fuzzy",
                    }
        return None
    except Exception as exc:
        LOGGER.warning("catalog ID lookup failed: %s", exc)
        return None


def find_top_candidates(
    product_name,
    *,
    raw_search_tokens,
    preferred_brand_token,
    strict_candidate_allowed,
    calc_name_match,
    priority_model_queries=None,
    b2b_search_candidates=None,
    top_n=5,
    min_score=0.40,
    allow_b2b=True,
):
    name = str(product_name or "").strip()
    if not name:
        return []
    try:
        with db_connection() as conn:
            brand = preferred_brand_token(name)
            candidates = {}

            _add_fts_candidate_rows(
                candidates,
                conn,
                name,
                brand=brand,
                priority_model_queries=priority_model_queries,
            )

            if len(candidates) < 40 and brand and len(brand) >= 3:
                _add_candidate_rows(candidates, conn.execute(
                    "SELECT ni.onliner_id, ni.raw_name, oc.url "
                    "FROM name_index ni JOIN onliner_catalog oc ON oc.onliner_id = ni.onliner_id "
                    "WHERE ni.raw_name LIKE ? LIMIT 200",
                    (f"%{brand}%",),
                ).fetchall())

            for token in raw_search_tokens(name)[:3]:
                if len(candidates) >= 160:
                    break
                if len(token) >= 4:
                    _add_candidate_rows(candidates, conn.execute(
                        "SELECT ni.onliner_id, ni.raw_name, oc.url "
                        "FROM name_index ni JOIN onliner_catalog oc ON oc.onliner_id = ni.onliner_id "
                        "WHERE ni.raw_name LIKE ? LIMIT 100",
                        (f"%{token}%",),
                    ).fetchall())

        if not candidates:
            return []

        scored = []
        for oid, (candidate_name, candidate_url) in candidates.items():
            allowed, _reason = strict_candidate_allowed(name, candidate_name)
            if not allowed:
                continue
            match = _best_candidate_match(calc_name_match, name, candidate_name, candidate_url)
            score = float(match.get("score", 0.0) or 0.0)
            if bool(match.get("match", False)) and score >= min_score:
                scored.append({
                    "id": oid,
                    "name": candidate_name,
                    "url": candidate_url,
                    "score": round(score, 3),
                    "source": "local_db",
                    "reason": str(match.get("reason", "") or ""),
                })

        scored.sort(key=lambda item: item["score"], reverse=True)
        top_scored = scored[:top_n]
        seen_ids = {normalize_onliner_id(item.get("id", "")) for item in top_scored}
        if allow_b2b and b2b_search_candidates and len(top_scored) < top_n:
            for candidate in b2b_search_candidates(name, category_name=infer_category(name), limit=top_n * 2):
                candidate_id = normalize_onliner_id(candidate.get("id", ""))
                score = float(candidate.get("score", 0.0) or 0.0)
                if not candidate_id or candidate_id in seen_ids or score < min_score:
                    continue
                top_scored.append({
                    "id": candidate_id,
                    "name": str(candidate.get("name", "") or "").strip(),
                    "url": str(candidate.get("url", "") or "").strip(),
                    "score": round(score, 3),
                    "source": "b2b_fuzzy",
                })
                seen_ids.add(candidate_id)
                if len(top_scored) >= top_n:
                    break
        top_scored.sort(key=lambda item: float(item.get("score", 0.0) or 0.0), reverse=True)
        return top_scored[:top_n]
    except Exception as exc:
        LOGGER.warning("catalog candidate lookup failed: %s", exc)
        return []


def search_tgpc_pc_candidates(
    local_name,
    *,
    find_exact,
    tgpc_pc_code_queries,
    extract_tgpc_pc_code,
    is_tgpc_pc_name,
    calc_name_match,
    limit=12,
):
    name = str(local_name or "").strip()
    if not name or not is_tgpc_pc_name(name):
        return []
    limit = max(5, min(int(limit or 12), 80))

    pool = {}
    exact = find_exact(name)
    if isinstance(exact, dict):
        exact_id = normalize_onliner_id(exact.get("id", ""))
        if exact_id:
            pool[exact_id] = (
                str(exact.get("name", "") or "").strip(),
                str(exact.get("url", "") or "").strip(),
            )

    try:
        with db_connection() as conn:
            queries = list(tgpc_pc_code_queries(name))
            code = extract_tgpc_pc_code(name)
            fts_available = False
            if code and len(code) >= 4:
                try:
                    rows = conn.execute(
                        "SELECT f.onliner_id, f.raw_name, oc.url "
                        "FROM name_index_fts f "
                        "LEFT JOIN onliner_catalog oc ON oc.onliner_id = f.onliner_id "
                        "WHERE name_index_fts MATCH ? LIMIT 220",
                        (f'"{code}"',),
                    ).fetchall()
                    fts_available = True
                except Exception:
                    rows = []
                for row in rows:
                    oid = normalize_onliner_id(row["onliner_id"])
                    raw = str(row["raw_name"] or "").strip()
                    url = str(row["url"] or "").strip()
                    if oid and raw:
                        pool[oid] = (raw, url)
            if not fts_available:
                if code and len(code) >= 4:
                    queries.append(code)
                for query in queries:
                    query = str(query or "").strip()
                    if len(query) < 2:
                        continue
                    rows = conn.execute(
                        "SELECT ni.onliner_id, ni.raw_name, oc.url "
                        "FROM name_index ni JOIN onliner_catalog oc ON oc.onliner_id = ni.onliner_id "
                        "WHERE ni.raw_name LIKE ? LIMIT 220",
                        (f"%{query}%",),
                    ).fetchall()
                    for row in rows:
                        oid = normalize_onliner_id(row["onliner_id"])
                        raw = str(row["raw_name"] or "").strip()
                        url = str(row["url"] or "").strip()
                        if oid and raw:
                            pool[oid] = (raw, url)
    except Exception as exc:
        LOGGER.warning("catalog TGPC candidate lookup failed: %s", exc)
        return []

    local_tgpc_code = extract_tgpc_pc_code(name)
    items = []
    seen = set()
    for oid, (candidate_name, candidate_url) in pool.items():
        if not oid or not candidate_name or oid in seen:
            continue
        candidate_lower = candidate_name.lower()
        url_lower = str(candidate_url or "").strip().lower()
        candidate_is_tgpc = "tgpc" in candidate_lower
        candidate_is_pc_url = bool(url_lower and any(part in url_lower for part in ("/desktop/", "/computer/", "/tgpc/")))
        if not candidate_is_tgpc and not candidate_is_pc_url:
            continue

        match = calc_name_match(name, candidate_name)
        score = float(match.get("score", 0.0) or 0.0)
        if local_tgpc_code:
            candidate_code = extract_tgpc_pc_code(candidate_name)
            if candidate_code and candidate_code != local_tgpc_code:
                score = min(score, 0.12)
            elif not candidate_code and local_tgpc_code not in (candidate_name + str(candidate_url or "")):
                score = min(score, 0.68)
        if not candidate_is_tgpc:
            score *= 0.70
        if url_lower and not any(part in url_lower for part in ("/desktop/", "/computer/", "/tgpc/")):
            score *= 0.80

        if score < 0.34:
            continue
        seen.add(oid)
        items.append({
            "id": oid,
            "name": candidate_name,
            "url": str(candidate_url or "").strip(),
            "score": round(score, 3),
            "source": "db_tgpc_pc",
            "reason": str(match.get("reason", "") or ""),
        })

    items.sort(key=lambda item: float(item.get("score", 0.0) or 0.0), reverse=True)
    return items[:limit]



ONLINER_FTS_READY = False
ONLINER_FTS_READY_PATH = ""


def _ensure_search_fts(conn, force=False):
    """Build a fast FTS index for read-only CRM lookups."""
    global ONLINER_FTS_READY, ONLINER_FTS_READY_PATH
    db_rows = conn.execute("PRAGMA database_list").fetchall()
    db_path = str(db_rows[0][2] or "") if db_rows else ""
    if ONLINER_FTS_READY and ONLINER_FTS_READY_PATH == db_path and not force:
        return True
    try:
        conn.execute("""
            CREATE VIRTUAL TABLE IF NOT EXISTS name_index_fts USING fts5(
                raw_name,
                name_key,
                onliner_id UNINDEXED,
                source UNINDEXED,
                tokenize='unicode61'
            )
        """)
        conn.executescript("""
            CREATE TRIGGER IF NOT EXISTS name_index_fts_after_insert
            AFTER INSERT ON name_index BEGIN
                INSERT INTO name_index_fts(raw_name, name_key, onliner_id, source)
                VALUES (new.raw_name, new.name_key, new.onliner_id, new.source);
            END;
            CREATE TRIGGER IF NOT EXISTS name_index_fts_after_delete
            AFTER DELETE ON name_index BEGIN
                DELETE FROM name_index_fts WHERE name_key = old.name_key;
            END;
            CREATE TRIGGER IF NOT EXISTS name_index_fts_after_update
            AFTER UPDATE ON name_index BEGIN
                DELETE FROM name_index_fts WHERE name_key = old.name_key;
                INSERT INTO name_index_fts(raw_name, name_key, onliner_id, source)
                VALUES (new.raw_name, new.name_key, new.onliner_id, new.source);
            END;
        """)
        source_count = int(conn.execute("SELECT COUNT(*) FROM name_index").fetchone()[0] or 0)
        fts_count = int(conn.execute("SELECT COUNT(*) FROM name_index_fts").fetchone()[0] or 0)
        if force or source_count != fts_count:
            conn.execute("DELETE FROM name_index_fts")
            conn.execute("""
                INSERT INTO name_index_fts(raw_name, name_key, onliner_id, source)
                SELECT raw_name, name_key, onliner_id, source FROM name_index
            """)
        ONLINER_FTS_READY = True
        ONLINER_FTS_READY_PATH = db_path
        return True
    except Exception as exc:
        LOGGER.warning("catalog FTS index unavailable: %s", exc)
        return False


def _fts_search_tokens(query):
    text = str(query or "").strip().lower()
    # Match the tokenizer used by SQLite FTS5. Keeping hyphens inside a token
    # turned `IS-47-XT` into `is47xt`, which can never match FTS terms
    # `is`, `47`, `xt` and caused exact articles to disappear from candidates.
    tokens = re.findall(r"[0-9a-zа-яё]{2,}", text, flags=re.I)
    cleaned = []
    seen = set()
    stop = {
        "lcd", "led", "oled", "ips", "va", "tn", "hdr", "usb",
        "монитор", "игровой", "товар", "ноутбук", "компьютер", "пэвм",
        "для", "шт", "черный", "белый", "black", "white",
    }
    for token in tokens:
        token = re.sub(r"^[^0-9a-zа-яё]+|[^0-9a-zа-яё]+$", "", token.lower(), flags=re.I)
        if len(token) < 2 or token in stop:
            continue
        if token not in seen:
            seen.add(token)
            cleaned.append(token)
    return cleaned[:10]


def _add_fts_candidate_rows(target, conn, query, *, brand="", priority_model_queries=None):
    if not _ensure_search_fts(conn):
        return

    attempts = []
    seen_attempts = set()

    def add_attempt(tokens, operator="AND", limit=120):
        tokens = [str(token or "").strip() for token in tokens if str(token or "").strip()]
        expr = _fts_expr(tokens[:8], operator)
        if not expr or expr in seen_attempts:
            return
        seen_attempts.add(expr)
        attempts.append((expr, int(limit)))

    priority_queries = []
    if callable(priority_model_queries):
        try:
            priority_queries = list(priority_model_queries(query) or [])
        except Exception:
            priority_queries = []
    for priority_query in priority_queries[:6]:
        priority_tokens = _fts_search_tokens(priority_query)
        if priority_tokens:
            add_attempt(priority_tokens, limit=100)

    query_tokens = _fts_search_tokens(query)
    model_tokens = [
        token
        for token in query_tokens
        if any(ch.isdigit() for ch in token) and any(ch.isalpha() for ch in token)
    ]
    if model_tokens:
        add_attempt(model_tokens[:6], limit=140)

    brand_tokens = _fts_search_tokens(brand)
    if brand_tokens:
        for start in range(len(query_tokens)):
            if query_tokens[start : start + len(brand_tokens)] != brand_tokens:
                continue
            add_attempt(query_tokens[start : start + len(brand_tokens) + 4], limit=120)
            break

    for model_token in model_tokens[:4]:
        try:
            model_idx = query_tokens.index(model_token)
        except ValueError:
            continue
        add_attempt(query_tokens[max(0, model_idx - 3) : model_idx + 2], limit=100)

    if query_tokens:
        add_attempt(query_tokens[:7], limit=100)
        add_attempt(query_tokens[:8], operator="OR", limit=160)

    for expr, limit in attempts[:14]:
        try:
            rows = conn.execute(
                "SELECT f.onliner_id, f.raw_name, oc.url "
                "FROM name_index_fts f "
                "LEFT JOIN onliner_catalog oc ON oc.onliner_id = f.onliner_id "
                "WHERE name_index_fts MATCH ? "
                "ORDER BY rank LIMIT ?",
                (expr, limit),
            ).fetchall()
        except Exception:
            continue
        _add_candidate_rows(target, rows)


def _fts_term(token):
    token = re.sub(r"[^0-9a-zа-яё]+", "", str(token or "").lower(), flags=re.I)
    if not token:
        return ""
    return '"' + token.replace('"', '""') + '"' + ("*" if len(token) >= 3 else "")


def _fts_expr(tokens, operator="AND"):
    parts = [_fts_term(t) for t in tokens]
    parts = [p for p in parts if p]
    if not parts:
        return ""
    joiner = " OR " if operator == "OR" else " "
    return joiner.join(parts)


def _score_search_item(query_lower, query_tokens, raw_name):
    name = str(raw_name or "").lower()
    if not name:
        return 0
    score = 0
    if query_lower and query_lower in name:
        score += 1000
    for token in query_tokens:
        if token in name:
            score += 80 + min(len(token), 20)
        else:
            compact_name = re.sub(r"[^0-9a-zа-яё]+", "", name, flags=re.I)
            compact_token = re.sub(r"[^0-9a-zа-яё]+", "", token, flags=re.I)
            if compact_token and compact_token in compact_name:
                score += 60 + min(len(compact_token), 20)
    score -= min(len(name), 250) / 1000.0
    return score


def stats_payload():
    try:
        with db_connection() as conn:
            total_products = conn.execute("SELECT COUNT(*) FROM onliner_catalog").fetchone()[0]
            total_names = conn.execute("SELECT COUNT(*) FROM name_index").fetchone()[0]
            by_source = {
                row[0]: row[1]
                for row in conn.execute(
                    "SELECT source, COUNT(*) FROM onliner_catalog GROUP BY source ORDER BY COUNT(*) DESC"
                ).fetchall()
            }
        return {"total_products": total_products, "total_names": total_names, "by_source": by_source}
    except Exception:
        return {"total_products": 0, "total_names": 0, "by_source": {}}


def search_payload(query):
    query = str(query or "").strip()
    if not query:
        return {"items": []}
    items = []
    seen = set()
    query_lower = query.lower()
    compact = re.sub(r"[^a-z0-9а-яё]+", "", query_lower, flags=re.I)
    tokens = _fts_search_tokens(query)
    model_tokens = [t for t in tokens if re.search(r"[0-9]", t) and re.search(r"[a-zа-яё]", t, flags=re.I)]

    def add_rows(rows):
        for row in rows:
            keys = row.keys()
            oid = normalize_onliner_id(row["onliner_id"])
            raw_name = str((row["raw_name"] if "raw_name" in keys else row["name"]) or "").strip()
            if not oid or not raw_name or oid in seen:
                continue
            seen.add(oid)
            items.append({
                "id": oid,
                "name": raw_name,
                "url": str(row["url"] or "").strip() if "url" in keys else "",
                "source": str(row["source"] or "").strip() if "source" in keys else "",
                "_score": _score_search_item(query_lower, tokens, raw_name),
            })

    try:
        with db_connection() as conn:
            if query.isdigit():
                add_rows(conn.execute(
                    "SELECT oc.onliner_id, oc.name, oc.url, oc.source "
                    "FROM onliner_catalog oc WHERE oc.onliner_id = ? LIMIT 20",
                    (query,),
                ).fetchall())

            used_fts = _ensure_search_fts(conn)
            if used_fts:
                fts_attempts = []
                if model_tokens:
                    fts_attempts.append(_fts_expr(model_tokens, "AND"))
                    for token in model_tokens[:3]:
                        fts_attempts.append(_fts_expr([token], "AND"))
                if tokens:
                    meaningful = [t for t in tokens if t not in model_tokens][:4]
                    if model_tokens and meaningful:
                        fts_attempts.append(_fts_expr(meaningful[:2] + model_tokens[:2], "AND"))
                    fts_attempts.append(_fts_expr(tokens[:5], "AND"))
                    fts_attempts.append(_fts_expr(tokens[:6], "OR"))

                tried = set()
                for expr in fts_attempts:
                    if not expr or expr in tried or len(items) >= 80:
                        continue
                    tried.add(expr)
                    add_rows(conn.execute(
                        "SELECT f.onliner_id, f.raw_name, oc.url, COALESCE(oc.source, f.source) AS source "
                        "FROM name_index_fts f "
                        "LEFT JOIN onliner_catalog oc ON oc.onliner_id = f.onliner_id "
                        "WHERE name_index_fts MATCH ? "
                        "ORDER BY rank LIMIT 80",
                        (expr,),
                    ).fetchall())

            if not items:
                rows = conn.execute(
                    "SELECT ni.onliner_id, ni.raw_name, oc.url, oc.source "
                    "FROM name_index ni "
                    "LEFT JOIN onliner_catalog oc ON oc.onliner_id = ni.onliner_id "
                    "WHERE lower(ni.raw_name) LIKE ? LIMIT 40",
                    (f"%{query_lower}%",),
                ).fetchall()
                add_rows(rows)

            if not items and compact and compact != query_lower:
                rows = conn.execute(
                    "SELECT ni.onliner_id, ni.raw_name, oc.url, oc.source "
                    "FROM name_index ni "
                    "LEFT JOIN onliner_catalog oc ON oc.onliner_id = ni.onliner_id "
                    "WHERE replace(replace(replace(lower(ni.raw_name), '-', ''), ' ', ''), '\"', '') LIKE ? LIMIT 40",
                    (f"%{compact}%",),
                ).fetchall()
                add_rows(rows)
    except Exception as exc:
        return {"items": [], "message": str(exc)}, 500

    items.sort(key=lambda item: (
        -float(item.get("_score", 0) or 0),
        0 if query_lower in str(item.get("name", "")).lower() else 1,
        len(str(item.get("name", ""))),
        str(item.get("name", "")).lower(),
    ))
    for item in items:
        item.pop("_score", None)
    return {"items": items[:30]}


def rebuild_payload(session_dir, *, read_consolidated_df, populate_from_dataframe):
    if not session_dir:
        return {"status": "error", "message": "Нет активной сессии"}, 400
    try:
        df = read_consolidated_df(session_dir)
        products, names = populate_from_dataframe(df, "price_load", skip_suppliers=["N-Tech", "TGPC"])
        return {"status": "ok", "products": products, "names": names}
    except Exception as exc:
        return {"status": "error", "message": str(exc)}, 500


def _reset_import_status(message):
    with CATALOG_IMPORT_LOCK:
        catalog_import_status.update({
            "job_id": new_job_id(),
            "running": True,
            "total": 0,
            "done": 0,
            "inserted": 0,
            "skipped": 0,
            "existing": 0,
            "message": message,
            "percent": 0,
            "finished_at": None,
        })


def _catalog_rows(filepath, file_ext):
    if file_ext in (".xlsx", ".xls"):
        import openpyxl

        workbook = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        worksheet = workbook.active
        return [[str(cell.value or "") for cell in row] for row in worksheet.iter_rows()]
    with open(filepath, encoding="utf-8-sig", newline="") as handle:
        return [[str(cell) for cell in row] for row in csv.reader(handle)]


def catalog_import_worker(
    filepath: str,
    file_ext: str,
    *,
    normalize_name_key,
    cleanup_file: bool = True,
):
    _reset_import_status("Загружаю список существующих товаров...")
    job_id = str(catalog_import_status.get("job_id", "") or "")
    with log_context(job_id=job_id):
        return _catalog_import_worker(
            filepath,
            file_ext,
            normalize_name_key=normalize_name_key,
            cleanup_file=cleanup_file,
        )


def _catalog_import_worker(
    filepath: str,
    file_ext: str,
    *,
    normalize_name_key,
    cleanup_file: bool,
):
    try:
        existing_ids = set()
        try:
            with DB_WRITE_LOCK:
                with db_connection() as conn:
                    existing_ids = {row[0] for row in conn.execute("SELECT onliner_id FROM onliner_catalog").fetchall()}
        except Exception as exc:
            LOGGER.warning("catalog import could not preload existing IDs: %s", exc)
            if _is_corrupt_db_error(exc):
                with DB_WRITE_LOCK:
                    quarantine_onliner_db()
                    init_onliner_db()
                existing_ids = set()

        existing_count = len(existing_ids)
        LOGGER.info("catalog import started existing_products=%s", existing_count)
        with CATALOG_IMPORT_LOCK:
            catalog_import_status["message"] = f"В базе {existing_count:,} товаров. Читаю файл..."

        rows = _catalog_rows(filepath, file_ext)
        if rows:
            rows = rows[1:]
        total = len(rows)
        with CATALOG_IMPORT_LOCK:
            catalog_import_status["total"] = total
            catalog_import_status["message"] = f"В файле {total:,} строк. Фильтрую новые..."

        inserted = 0
        categories_updated = 0
        skipped = 0
        already_existing = 0
        now_ts = int(time.time())
        batch_size = 2000

        for batch_start in range(0, total, batch_size):
            batch = rows[batch_start: batch_start + batch_size]
            product_rows = []
            name_rows = []
            new_product_ids = set()
            for row in batch:
                while len(row) < 8:
                    row.append("")
                category = normalize_catalog_category_name(row[0].strip())
                model_short = row[2].strip()
                oid = normalize_onliner_id(row[4].strip())
                full_name = row[7].strip()
                if not oid:
                    skipped += 1
                    continue
                is_existing = oid in existing_ids
                if is_existing:
                    already_existing += 1

                name = full_name if full_name else (f"{category} {model_short}".strip())
                if not name:
                    skipped += 1
                    continue
                product_rows.append((oid, name, "", category, "onliner_catalog", now_ts))
                if is_existing and category:
                    categories_updated += 1
                elif not is_existing:
                    new_product_ids.add(oid)

                name_key = normalize_name_key(name)
                if name_key:
                    name_rows.append((name_key, oid, name))
                if model_short and model_short.lower() not in name.lower():
                    alt_name = f"{category} {model_short}".strip() if category else model_short
                    alt_key = normalize_name_key(alt_name)
                    if alt_key and alt_key != name_key:
                        name_rows.append((alt_key, oid, alt_name))

            try:
                if product_rows:
                    with DB_WRITE_LOCK:
                        with db_connection() as conn:
                            conn.executemany(
                                "INSERT INTO onliner_catalog "
                                "(onliner_id, name, url, category, source, updated_at) VALUES (?,?,?,?,?,?) "
                                "ON CONFLICT(onliner_id) DO UPDATE SET "
                                "name=excluded.name, "
                                "category=CASE WHEN excluded.category <> '' THEN excluded.category ELSE onliner_catalog.category END, "
                                "source=excluded.source, updated_at=excluded.updated_at",
                                product_rows,
                            )
                            conn.executemany(
                                "INSERT OR IGNORE INTO name_index "
                                "(name_key, onliner_id, raw_name) VALUES (?,?,?)",
                                name_rows,
                            )
                    for product_row in product_rows:
                        existing_ids.add(product_row[0])
                inserted += len(new_product_ids)
            except Exception as exc:
                LOGGER.warning("catalog import batch failed: %s", exc)
                if _is_corrupt_db_error(exc):
                    with DB_WRITE_LOCK:
                        quarantine_onliner_db()
                        init_onliner_db()
                    existing_ids = set()
                    try:
                        if product_rows:
                            with DB_WRITE_LOCK:
                                with db_connection() as conn:
                                    conn.executemany(
                                        "INSERT INTO onliner_catalog "
                                        "(onliner_id, name, url, category, source, updated_at) VALUES (?,?,?,?,?,?) "
                                        "ON CONFLICT(onliner_id) DO UPDATE SET "
                                        "name=excluded.name, "
                                        "category=CASE WHEN excluded.category <> '' THEN excluded.category ELSE onliner_catalog.category END, "
                                        "source=excluded.source, updated_at=excluded.updated_at",
                                        product_rows,
                                    )
                                    conn.executemany(
                                        "INSERT OR IGNORE INTO name_index "
                                        "(name_key, onliner_id, raw_name) VALUES (?,?,?)",
                                        name_rows,
                                    )
                            for product_row in product_rows:
                                existing_ids.add(product_row[0])
                            inserted += len(new_product_ids)
                    except Exception as retry_exc:
                        LOGGER.error("catalog import batch retry failed: %s", retry_exc)

            done = batch_start + len(batch)
            percent = round(done / total * 100) if total else 100
            with CATALOG_IMPORT_LOCK:
                catalog_import_status.update({
                    "done": done,
                    "inserted": inserted,
                    "skipped": skipped,
                    "existing": already_existing,
                    "percent": percent,
                    "message": (
                        f"Обработано {done:,} из {total:,} | Новых: {inserted:,} "
                        f"| Обновлено категорий: {categories_updated:,} "
                        f"| Уже было: {already_existing:,} | Пропущено: {skipped:,}"
                    ),
                })

        new_total = existing_count + inserted
        try:
            with DB_WRITE_LOCK:
                with db_connection() as conn:
                    new_total = int(conn.execute("SELECT COUNT(*) FROM onliner_catalog").fetchone()[0] or 0)
        except Exception:
            pass
        with CATALOG_IMPORT_LOCK:
            catalog_import_status.update({
                "running": False,
                "done": total,
                "inserted": inserted,
                "skipped": skipped,
                "existing": already_existing,
                "percent": 100,
                "message": (
                    f"Готово! Всего в базе: {new_total:,} (+{inserted:,} новых, "
                    f"{already_existing:,} уже было, {categories_updated:,} категорий обновлено, "
                    f"{skipped:,} пропущено)"
                ),
                "finished_at": int(time.time()),
            })
        LOGGER.info(
            "catalog import completed existing_before=%s new_inserted=%s "
            "categories_updated=%s skipped=%s already_existing=%s total_now=%s",
            existing_count,
            inserted,
            categories_updated,
            skipped,
            already_existing,
            new_total,
        )
    except Exception as exc:
        with CATALOG_IMPORT_LOCK:
            catalog_import_status.update({
                "running": False,
                "message": f"Ошибка: {exc}",
                "finished_at": int(time.time()),
            })
        LOGGER.exception("catalog import failed")
    finally:
        if cleanup_file:
            try:
                os.remove(filepath)
            except Exception:
                pass


def import_gsheet_payload(payload, *, normalize_name_key, start_thread: Callable[[Callable[[], None]], None]):
    with CATALOG_IMPORT_LOCK:
        if catalog_import_status.get("running"):
            return {"status": "already_running", "message": "Импорт уже запущен"}, 409

    payload = payload if isinstance(payload, dict) else {}
    sheet_id = str(payload.get("sheet_id", "")).strip()
    sheet_name = str(payload.get("sheet_name", "All_Catalog")).strip() or "All_Catalog"
    force_refresh = bool(payload.get("force_refresh", False))
    if not sheet_id:
        return {"status": "error", "message": "sheet_id required"}, 400

    def _download_and_import():
        urls = [
            (
                f"https://docs.google.com/spreadsheets/d/{sheet_id}"
                f"/export?format=csv&sheet={urllib.parse.quote(sheet_name)}"
            ),
            (
                f"https://docs.google.com/spreadsheets/d/{sheet_id}"
                f"/gviz/tq?tqx=out:csv&sheet={urllib.parse.quote(sheet_name)}"
            ),
        ]
        safe_sheet = re.sub(r"[^a-zA-Z0-9._-]+", "_", sheet_name).strip("._-") or "sheet"
        cache_path = ONLINER_DB_GSHEET_CACHE_DIR / f"{sheet_id}_{safe_sheet}.csv"
        with CATALOG_IMPORT_LOCK:
            catalog_import_status.update({
                "running": True,
                "total": 0,
                "done": 0,
                "inserted": 0,
                "skipped": 0,
                "percent": 0,
                "finished_at": None,
                "message": "Подключаюсь к Google Sheets…",
            })
        try:
            ONLINER_DB_GSHEET_CACHE_DIR.mkdir(parents=True, exist_ok=True)
            now_ts = int(time.time())
            cache_fresh = False
            if cache_path.exists():
                try:
                    age_sec = max(0, now_ts - int(cache_path.stat().st_mtime))
                    cache_fresh = (age_sec <= ONLINER_DB_GSHEET_CACHE_TTL_SEC) and (cache_path.stat().st_size > 0)
                except Exception:
                    cache_fresh = False

            if cache_fresh and not force_refresh:
                with CATALOG_IMPORT_LOCK:
                    catalog_import_status.update({
                        "message": "Использую локальный кэш Google Sheets (без повторного скачивания)…",
                        "percent": 5,
                    })
            else:
                tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".csv")
                tmp.close()
                downloaded = False
                last_error = None
                for url in urls:
                    try:
                        request = urllib.request.Request(
                            url,
                            headers={"User-Agent": "Mozilla/5.0 (compatible; PriceMixer/1.0)"},
                        )
                        chunk_size = 131072
                        with urllib.request.urlopen(request, timeout=300) as response:
                            total_size = int(response.headers.get("Content-Length") or 0)
                            received = 0
                            with open(tmp.name, "wb") as handle:
                                while True:
                                    chunk = response.read(chunk_size)
                                    if not chunk:
                                        break
                                    handle.write(chunk)
                                    received += len(chunk)
                                    mb = received / 1024 / 1024
                                    total_mb = total_size / 1024 / 1024 if total_size else 0
                                    message = (
                                        f"Скачано {mb:.1f} / {total_mb:.1f} МБ…"
                                        if total_size else f"Скачано {mb:.1f} МБ…"
                                    )
                                    percent = int(received / total_size * 15) if total_size else 5
                                    with CATALOG_IMPORT_LOCK:
                                        catalog_import_status.update({"message": message, "percent": percent})
                        downloaded = True
                        break
                    except Exception as exc:
                        last_error = exc
                        continue
                if not downloaded:
                    raise RuntimeError(f"Не удалось скачать файл: {last_error}")
                try:
                    shutil.move(tmp.name, str(cache_path))
                except Exception:
                    cache_path = Path(tmp.name)
                with CATALOG_IMPORT_LOCK:
                    catalog_import_status.update({
                        "message": "Файл скачан и сохранен в кэш. Начинаю импорт…",
                        "percent": 16,
                    })

            catalog_import_worker(str(cache_path), ".csv", normalize_name_key=normalize_name_key, cleanup_file=False)
        except Exception as exc:
            with CATALOG_IMPORT_LOCK:
                catalog_import_status.update({
                    "running": False,
                    "message": f"Ошибка скачивания: {exc}",
                    "finished_at": int(time.time()),
                })

    start_thread(_download_and_import)
    return {"status": "started", "message": "Скачиваю и импортирую…"}


def import_csv_payload(file, *, normalize_name_key, start_thread: Callable[[Callable[[], None], str, str], None]):
    with CATALOG_IMPORT_LOCK:
        if catalog_import_status.get("running"):
            return {"status": "already_running", "message": "Импорт уже запущен"}, 409

    if not file or not file.filename:
        return {"status": "error", "message": "Файл не передан"}, 400

    filename = file.filename.lower()
    if filename.endswith(".xlsx"):
        ext = ".xlsx"
    elif filename.endswith(".xls"):
        ext = ".xls"
    elif filename.endswith(".csv"):
        ext = ".csv"
    else:
        return {"status": "error", "message": "Поддерживаются только CSV и XLSX"}, 400

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=ext)
    file.save(tmp.name)
    tmp.close()
    start_thread(lambda: catalog_import_worker(tmp.name, ext, normalize_name_key=normalize_name_key))
    return {"status": "started", "message": "Импорт запущен"}


def import_status_payload():
    with CATALOG_IMPORT_LOCK:
        return dict(catalog_import_status)
